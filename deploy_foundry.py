#!/usr/bin/env python3
"""Azure AI Foundry 统一部署入口（中文注释版，Python 3.10+）。

参考原 AIFoundry_az_logen.py 和 deploy_foundry_models-wuhao01.py 的部署流程；
不导入旧脚本或其密钥。本命令行入口默认只预览，只有 --apply 才访问 Azure。
可视化入口通过本地服务提交临时计划与一次性凭据，只有明确开始部署才执行 --apply。
不会删除资源或回收其他部署的配额；内容防护开关可以选择 High 或 Low 阈值。
JSON/CSV 部署和格式化 Excel 导出只需标准库；读取 Excel 计划需要 openpyxl。详细说明见 README.md。
"""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from queue import Queue
from threading import RLock
import argparse
import csv
import hashlib
import io
import json
import math
import os
import random
import re
import shutil
import socket
import ssl
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from dataclasses import asdict, dataclass, fields, replace
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from foundry_regions import normalize_location
from foundry_excel import write_excel_report

# 按微软官方 azure-rest-api-specs 核对的稳定 API，不盲目使用 preview。
ARM = "https://management.azure.com"
COGNITIVE_API = "2025-06-01"
# 配额 usages 接口使用 Microsoft 文档明确列出的稳定 API 版本，避免与资源创建 API 混用。
USAGE_API = "2024-10-01"
RESOURCE_API = "2021-04-01"
PROVIDER_API = "2021-04-01"
DEFAULT_SKU = "GlobalStandard"
# 控制台两档防护：均开启阻止，High 是无需特殊审批的最低常规阈值。
GUARD_POLICIES = {"foundry-console-guard-high-v1": "High", "foundry-console-guard-low-v1": "Low"}
# 本入口仅部署按量 SKU，不会意外创建 Provisioned 预留吞吐量部署。
SUPPORTED_SKUS = {"GlobalStandard", "Standard", "DataZoneStandard"}
RETRY_STATUSES = {429, 500, 502, 503, 504}
FAILED_STATES = {"failed", "canceled", "cancelled", "deleted"}
_MISSING = object()
_REDACT_VALUES: set[str] = set()
CREDENTIAL_KEYS = {"AZURE_TENANT_ID", "AZURE_CLIENT_ID", "AZURE_CLIENT_SECRET"}
RUN_FLAGS = {"update_existing", "register_provider", "skip_quota_check"}
RUN_OPTIONS = RUN_FLAGS | {"export_api_key", "max_concurrency"}

# 兼容原工作簿的列头；星号、空格和大小写不影响识别。
ALIASES = {
    "subscription_id": ["Subscription Id", "SubscriptionID", "订阅ID", "订阅 Id"],
    "resource_group": ["Resource Group", "ResourceGroup", "资源组", "资源组名称"],
    "location": ["Foundry Region", "Global Standard Region", "Region", "Location", "区域", "位置"],
    "account_name": ["Foundry Resource Name", "Foundry Account Name", "AI Foundry Resource Name", "foundry_resource_name", "服务名称", "资源名称"],
    "project_name": ["Foundry Project Name", "AI Foundry Project Name", "Project Name", "项目名称"],
    "deployment_name": ["Deployment Name", "Foundry Deployment Name", "部署名称"],
    "model_name": ["Foundry Model", "Model Name", "模型名称", "(Azure OpenAI) Global Standard Model"],
    "model_version": ["Foundry Model Version", "Model Version", "模型版本"],
    "model_format": ["Foundry Model Format", "Model Format", "模型格式"],
    "sku_name": ["SKU Name", "Foundry SKU", "Deployment SKU", "SKU"],
    "capacity": ["Foundry Capacity", "Capacity", "容量(capacity)", "(Azure OpenAI) Global Standard Quota", "(Azure OpenAI) Global Standard Capacity", "Foundry Quota", "Quota"],
    "rai_policy_name": ["Content Filter", "RAI Policy Name", "内容筛选器"],
    "project_display_name": ["Project Display Name", "项目显示名称"],
    "project_description": ["Project Description", "项目描述"],
    "resource_group_mode": ["Resource Group Mode"],
    "account_mode": ["Account Mode"],
    "project_mode": ["Project Mode"],
}


class PlanError(ValueError):
    """计划错误：在创建任何云资源之前发现。"""


class DeployError(RuntimeError):
    """认证、配额、资源状态或部署失败。"""


class OperationUnknown(DeployError):
    """请求可能已生效，但未确认 Azure 最终状态；不得自动重新提交。"""


class OperationFailed(DeployError):
    """Azure 明确返回 Failed / Canceled 等终止状态。"""


class ArmError(DeployError):
    def __init__(self, status: int, code: str, message: str):
        self.status = status
        self.code = code
        super().__init__(f"HTTP {status} {code}: {message}")


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[\s_*]+", "", text(value)).casefold()


HEADER_MAP = {header_key(alias): key for key, aliases in ALIASES.items() for alias in [key, *aliases]}


def positive_int(value: Any, label: str = "capacity") -> int:
    """Excel 的 1.0 可接受；不把 1.5、布尔值、0 或负数悄悄截断。"""
    try:
        number = Decimal(str(value))
        if isinstance(value, bool) or not number.is_finite() or number != number.to_integral_value() or number <= 0:
            raise ValueError
        return int(number)
    except (InvalidOperation, ValueError):
        raise PlanError(f"{label} 必须是正整数") from None


@dataclass(frozen=True)
class DeployItem:
    subscription_id: str
    resource_group: str
    location: str
    account_name: str
    project_name: str
    deployment_name: str
    model_name: str
    model_version: str = ""
    model_format: str = "OpenAI"
    sku_name: str = DEFAULT_SKU
    capacity: int = 1
    project_display_name: str = ""
    project_description: str = ""
    rai_policy_name: str = ""
    resource_group_mode: str = "auto"
    account_mode: str = "auto"
    project_mode: str = "auto"

    @property
    def account_id(self) -> str:
        return f"/subscriptions/{self.subscription_id}/resourceGroups/{self.resource_group}/providers/Microsoft.CognitiveServices/accounts/{self.account_name}"

    @property
    def project_id(self) -> str:
        return f"{self.account_id}/projects/{self.project_name}"

    @property
    def deployment_id(self) -> str:
        # 部署属于 account，不是 projects 的子资源。
        return f"{self.account_id}/deployments/{self.deployment_name}"


def normalize_item(raw: dict[str, Any], row: int) -> DeployItem:
    def required(key: str) -> str:
        value = text(raw.get(key))
        if not value:
            raise PlanError(f"第 {row} 行缺少 {key}")
        return value

    try:
        subscription = str(uuid.UUID(required("subscription_id")))
    except ValueError:
        raise PlanError(f"第 {row} 行 subscription_id 必须是订阅 UUID，不是订阅显示名称") from None
    location = normalize_location(required("location"))
    if not re.fullmatch(r"[a-z][a-z0-9]{1,39}", location):
        raise PlanError(f"第 {row} 行区域无法识别，请使用 Azure location code")
    if location.startswith(("china", "usgov", "usdod")):
        raise PlanError("当前脚本只支持 Azure 公有云，不支持 China/Gov 区域")
    model = required("model_name")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", model):
        raise PlanError(f"第 {row} 行 model_name 含不支持的字符")
    group = text(raw.get("resource_group")) or f"rg-foundry-{location}"
    # 使用稳定哈希生成名称，同一计划重复执行不会随机创建新账号。
    digest = hashlib.sha256(f"{subscription}/{group.casefold()}/{location}".encode()).hexdigest()[:12]
    account = text(raw.get("account_name")) or f"fdry-{location}-{digest}"
    project = text(raw.get("project_name")) or "default-project"
    deployment = text(raw.get("deployment_name")) or model
    if not re.fullmatch(r"[\w().-]{1,90}", group) or group.endswith("."):
        raise PlanError(f"第 {row} 行资源组名称无效")
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,62}[a-z0-9]", account):
        raise PlanError(
            f"第 {row} 行 Foundry 资源名“{account}”无效：仅允许 2~64 位小写字母、数字和连字符（-），"
            "且首尾必须是字母或数字；留空可由系统自动生成。"
        )
    for label, name in [("project_name", project), ("deployment_name", deployment)]:
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", name) or (label == "project_name" and len(name) < 2):
            raise PlanError(f"第 {row} 行 {label} 应为字母、数字、点、下划线或连字符，项目名 2~64 位，部署名 1~64 位")
    sku = text(raw.get("sku_name")) or DEFAULT_SKU
    if sku not in SUPPORTED_SKUS:
        raise PlanError(f"第 {row} 行只支持按量 SKU：{', '.join(sorted(SUPPORTED_SKUS))}")
    version = text(raw.get("model_version"))
    if version and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", version):
        raise PlanError(f"第 {row} 行 model_version 无效；Excel 版本列请设置为文本")
    model_format = text(raw.get("model_format")) or "OpenAI"
    policy = text(raw.get("rai_policy_name"))
    if policy and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}", policy):
        raise PlanError(f"第 {row} 行内容筛选器名称无效")
    if policy and model_format.casefold() != "openai":
        raise PlanError(f"第 {row} 行仅 OpenAI 格式模型支持此内容筛选器选项；其他模型请使用 Azure 自身默认策略")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", model_format):
        raise PlanError(f"第 {row} 行 model_format 无效")
    modes = [text(raw.get(key)) or "auto" for key in ("resource_group_mode", "account_mode", "project_mode")]
    if any(value not in {"auto", "new", "existing"} for value in modes):
        raise PlanError(f"第 {row} 行资源来源无效")
    return DeployItem(subscription, group, location, account, project, deployment, model,
                      version, model_format, sku,
                      positive_int(raw.get("capacity") if text(raw.get("capacity")) else 1),
                      text(raw.get("project_display_name")), text(raw.get("project_description")), policy, *modes)


def table_records(rows: Iterable[Iterable[Any]]) -> Iterable[tuple[int, dict[str, Any]]]:
    """转换 Excel/CSV 行；忽略附加列，但拒绝重复列头和公式。"""
    iterator = iter(rows)
    headers = list(next(iterator, []))
    columns: dict[int, str] = {}
    for index, header in enumerate(headers):
        key = HEADER_MAP.get(header_key(header))
        if key:
            if key in columns.values():
                raise PlanError(f"重复列头: {key}")
            columns[index] = key
    missing = {"subscription_id", "location", "model_name"} - set(columns.values())
    if missing:
        raise PlanError(f"缺少必填列: {', '.join(sorted(missing))}")
    for number, row_values in enumerate(iterator, 2):
        row = list(row_values)
        raw = {key: row[index] if index < len(row) else None for index, key in columns.items()}
        if not any(text(v) for v in raw.values()):
            continue
        if text(raw.get("subscription_id")).startswith(("必填", "选填")):
            continue
        if any(isinstance(v, str) and v.startswith("=") for v in raw.values()):
            raise PlanError(f"第 {number} 行含公式；请粘贴为值后再部署")
        yield number, raw


def remember_secret(value: Any) -> None:
    """只在当前进程内登记脱敏值，不写环境变量、命令行或额外凭据文件。"""
    if isinstance(value, str) and value:
        _REDACT_VALUES.update({value, urllib.parse.quote(value, safe=""), urllib.parse.quote_plus(value)})


def read_json_document(path: Path) -> Any:
    """一次读入完整配置；重复键、语法错误不回显原文，避免泄露密钥。"""
    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise PlanError("JSON 存在重复字段，请删除重复项；为保护凭据不显示字段内容")
            result[key] = value
        return result

    try:
        document = json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=unique_object)
    except json.JSONDecodeError as exc:
        raise PlanError(f"JSON 格式错误：第 {exc.lineno} 行、第 {exc.colno} 列；请检查双引号、逗号，JSON 不支持注释") from None
    except UnicodeError:
        raise PlanError("JSON 文件编码无效，请保存为 UTF-8") from None
    # 先登记，再做其他验证；即使后续配置错误也不能把已知密钥写入报告。
    def collect(value):
        if isinstance(value, dict):
            for key, child in value.items():
                if str(key).casefold() in {"azure_client_secret", "client_secret", "access_token", "refresh_token"}:
                    remember_secret(child)
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)
    collect(document)
    return document


def validate_json_sections(document: dict) -> None:
    allowed = {"defaults", "deployments", "credentials", "run", "_comment"}
    if set(document) - allowed:
        raise PlanError("JSON 顶层只允许 defaults、deployments、credentials、run、_comment")
    if "credentials" in document:
        credentials = document["credentials"]
        if not isinstance(credentials, dict) or set(credentials) - CREDENTIAL_KEYS:
            raise PlanError("credentials 必须是对象，只允许三个 AZURE_ 凭据字段")
        if any(not isinstance(value, str) for value in credentials.values()):
            raise PlanError("credentials 的值必须是字符串")
        remember_secret(credentials.get("AZURE_CLIENT_SECRET"))
    if "run" in document:
        options = document["run"]
        if not isinstance(options, dict) or set(options) - RUN_OPTIONS:
            raise PlanError("run 只允许 update_existing、register_provider、skip_quota_check、export_api_key")
        if any(type(value) is not bool for key, value in options.items() if key != "max_concurrency"):
            raise PlanError("run 开关必须填写 true 或 false，不能加引号")
        workers = options.get("max_concurrency", 3)
        if type(workers) is not int or not 1 <= workers <= 8:
            raise PlanError("max_concurrency 必须为 1–8 的整数")


def load_plan(path: Path, sheet: str | None = None, *, document: Any = _MISSING) -> tuple[list[DeployItem], int]:
    """全量校验后才执行，避免第 N 行有错时前 N-1 行已创建资源。"""
    suffix = path.suffix.lower()
    if suffix == ".json":
        content = read_json_document(path) if document is _MISSING else document
        defaults: dict[str, Any] = {}
        if isinstance(content, dict):
            validate_json_sections(content)
            defaults = content.get("defaults", {})
            records = content.get("deployments")
        else:
            records = content
        allowed = {f.name for f in fields(DeployItem)}
        if not isinstance(defaults, dict) or set(defaults) - allowed:
            raise PlanError("JSON defaults 存在未知字段")
        if not isinstance(records, list):
            raise PlanError("JSON deployments 必须是列表")
        raw_rows = []
        for index, record in enumerate(records, 1):
            if not isinstance(record, dict) or set(record) - allowed:
                raise PlanError(f"JSON 第 {index} 项存在未知字段或不是对象")
            raw_rows.append((index, {**defaults, **record}))
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            raw_rows = list(table_records(csv.reader(stream)))
    elif suffix == ".xlsx":
        try:
            import openpyxl  # 可选依赖：JSON/CSV 和 --help 不依赖 Excel 库。
        except ImportError:
            raise PlanError("读取 Excel 需要 openpyxl，请运行 python -m pip install -r requirements.txt") from None
        try:
            book = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False)
        except (zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException):
            raise PlanError("Excel 文件损坏或不是有效的 .xlsx，请用 Excel 重新另存为 .xlsx 后重试") from None
        try:
            if sheet and sheet not in book.sheetnames:
                raise PlanError(f"工作表不存在: {sheet}")
            worksheet = book[sheet] if sheet else book.active
            raw_rows = list(table_records(worksheet.iter_rows(values_only=True)))
        finally:
            book.close()
    else:
        raise PlanError("部署计划仅支持 .json、.csv、.xlsx")
    plan: list[DeployItem] = []
    targets: dict[str, DeployItem] = {}
    accounts: dict[str, str] = {}
    projects: dict[str, tuple[str, str]] = {}
    duplicates = 0
    for number, raw in raw_rows:
        item = normalize_item(raw, number)
        account_key = item.account_id.casefold()
        if account_key in accounts and accounts[account_key] != item.location:
            raise PlanError(f"第 {number} 行同一 Foundry 账号指定了不同区域")
        accounts[account_key] = item.location
        project_info = (item.project_display_name, item.project_description)
        project_key = item.project_id.casefold()
        if project_key in projects and projects[project_key] != project_info:
            raise PlanError(f"第 {number} 行同一项目的显示名称/描述冲突")
        projects[project_key] = project_info
        target = item.deployment_id.casefold()
        if target in targets:
            if targets[target] != item:
                raise PlanError(f"第 {number} 行同一部署名配置冲突：{item.deployment_name}（部署名在账号内唯一）")
            duplicates += 1
            continue
        targets[target] = item
        plan.append(item)
    if not plan:
        raise PlanError("计划为空，没有可部署的数据行")
    return plan, duplicates

# -------------------- 认证与 ARM 通信层 --------------------
def redact(message: str, limit: int | None = 2000) -> str:
    """日志不记录令牌或客户端密钥；服务端错误仍需脱敏。"""
    secrets = _REDACT_VALUES | {os.environ.get(name, "") for name in ("AZURE_CLIENT_SECRET", "AZURE_ACCESS_TOKEN")}
    for secret in sorted((value for value in secrets if value), key=len, reverse=True):
        message = message.replace(secret, "[REDACTED]")
    message = re.sub(r"(?i)Bearer\s+[A-Za-z0-9._~+/-]+", "Bearer [REDACTED]", message)
    message = re.sub(r"(?i)(client_secret|access_token|refresh_token)([\s\"':=]+)[^\s,\"'}]+", r"\1\2[REDACTED]", message)
    return message if limit is None else message[:limit]


def token_http_error_message(status: int, response: bytes) -> str:
    """只提取 Entra 错误码；绝不回显可能含租户/应用信息的原始响应。"""
    try:
        payload = json.loads(response.decode("utf-8", errors="replace"))
    except (UnicodeError, ValueError):
        payload = {}
    if not isinstance(payload, dict):
        payload = {}
    codes = payload.get("error_codes")
    code = next((str(value) for value in codes if isinstance(value, int) and 10000 <= value <= 99999999), "") if isinstance(codes, list) else ""
    if not code:
        description = payload.get("error_description")
        match = re.search(r"AADSTS([0-9]{5,8})\b", description) if isinstance(description, str) else None
        code = match.group(1) if match else ""
    hints = {
        "7000215": "Client Secret 无效；请填写密钥的‘值’，不是‘密钥 ID’，并确认属于此应用",
        "7000222": "Client Secret 已过期；请创建新密钥并使用新密钥的‘值’",
        "700016": "应用在该租户中不存在；检查 APP_ID 与 TENANT_ID 是否来自同一租户",
        "90002": "租户不存在；检查 TENANT_ID",
        "7000218": "请求缺少有效的客户端密钥；检查 PASSWORD 是否已完整粘贴",
        "7000112": "应用已被禁用；请联系租户管理员",
    }
    if code:
        return f"Microsoft Entra 拒绝 SP 认证（AADSTS{code}）：{hints.get(code, '请让管理员根据错误码检查应用与凭据')}。未显示原始响应或密钥。"
    if status in {407, 502, 503, 504}:
        return f"Microsoft Entra 令牌请求返回 HTTP {status}；检查代理、防火墙及网络。"
    return f"Microsoft Entra 拒绝 SP 认证（HTTP {status}）；请核对租户、应用和密钥的值/有效期。未显示原始响应或密钥。"


def token_network_error_message(reason: object) -> str:
    """只返回预定义的连接故障类别，不泄露代理地址、请求体或 Secret。"""
    if isinstance(reason, ssl.SSLCertVerificationError):
        return "无法验证 Microsoft Entra 的 TLS 证书；请检查本机信任链、代理证书与 Python 环境。"
    if isinstance(reason, ssl.SSLError):
        return "连接 Microsoft Entra 时 TLS 握手失败；请检查代理、TLS 设置及 Python 环境。"
    if isinstance(reason, (socket.timeout, TimeoutError)):
        return "连接 Microsoft Entra 超时；请检查 Python 进程使用的代理与防火墙设置。"
    if isinstance(reason, socket.gaierror):
        return "无法解析 Microsoft Entra 域名；请检查本机 DNS 与 Python 代理设置。"
    if isinstance(reason, ConnectionRefusedError):
        return "连接 Microsoft Entra 被拒绝；请检查 Python 进程使用的代理与防火墙设置。"
    if isinstance(reason, str):
        if "407" in reason and ("proxy" in reason.lower() or "tunnel" in reason.lower()):
            return "代理要求身份验证（HTTP 407）；请检查 Python 进程的代理配置。"
        if "proxy" in reason.lower() or "tunnel" in reason.lower():
            return "连接 Microsoft Entra 的代理隧道失败；请检查 Python 进程的代理配置。"
    return "无法连接 Microsoft Entra 令牌服务；请检查本机 Python 的网络、代理、防火墙和 TLS 设置。"

class NoRedirect(urllib.request.HTTPRedirectHandler):
    """禁止自动重定向，防止 Authorization 被带到未经验证的地址。"""
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


class TokenProvider:
    def __init__(self, mode: str, credentials: dict[str, str] | None = None):
        self.mode = mode
        # JSON 凭据整组优先，不会把缺失字段与环境变量混用成另一套身份。
        self.credentials = dict(credentials) if credentials is not None else None
        if self.credentials is not None:
            remember_secret(self.credentials.get("AZURE_CLIENT_SECRET"))
        self.cache: dict[str, tuple[str, float]] = {}

    def get(self, subscription: str, refresh: bool = False) -> str:
        key = subscription if self.mode == "cli" else "service-principal"
        token, expiry = self.cache.get(key, ("", 0))
        if token and expiry > time.time() + 120 and not refresh:
            return token
        if self.mode == "cli":
            az = shutil.which("az") or shutil.which("az.cmd")
            if not az:
                raise DeployError("未找到 Azure CLI；请安装并 az login，或使用 --auth service-principal")
            # 每次显式指定订阅，不用 az account set 修改用户的全局 CLI 上下文。
            command = [az, "account", "get-access-token", "--subscription", subscription,
                       "--resource", ARM + "/", "--output", "json", "--only-show-errors"]
            try:
                completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8",
                                           errors="replace", timeout=90, check=False, shell=False)
            except (OSError, subprocess.TimeoutExpired):
                raise DeployError("Azure CLI 获取令牌失败或超时，请检查 az login 状态") from None
            if completed.returncode:
                # 不输出 CLI 的原始诊断数据，避免意外泄漏认证信息。
                raise DeployError("Azure CLI 获取令牌失败，请先 az login 并确认当前身份可访问目标订阅")
            try:
                result = json.loads(completed.stdout)
                token = result["accessToken"]
                expiry = float(result.get("expires_on") or time.time() + 300)
            except (ValueError, KeyError, TypeError):
                raise DeployError("Azure CLI 返回了无法解析的令牌响应") from None
        else:
            source = self.credentials if self.credentials is not None else os.environ
            tenant = source.get("AZURE_TENANT_ID", "")
            client = source.get("AZURE_CLIENT_ID", "")
            secret = source.get("AZURE_CLIENT_SECRET", "")
            try:
                tenant_id, client_id = uuid.UUID(tenant), uuid.UUID(client)
                if not tenant_id.int or not client_id.int:
                    raise ValueError
                tenant, client = str(tenant_id), str(client_id)
            except ValueError:
                raise DeployError("服务主体认证需要有效的 AZURE_TENANT_ID 和 AZURE_CLIENT_ID") from None
            if not secret or not secret.strip():
                raise DeployError("服务主体认证需要 AZURE_CLIENT_SECRET：请在 JSON credentials 或环境变量中填写密钥值")
            data = urllib.parse.urlencode({"grant_type": "client_credentials", "client_id": client,
                                           "client_secret": secret, "scope": ARM + "/.default"}).encode()
            request = urllib.request.Request(f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token", data=data)
            opener = urllib.request.build_opener(NoRedirect(), urllib.request.HTTPSHandler(context=ssl.create_default_context()))
            try:
                with opener.open(request, timeout=60) as response:
                    result = json.load(response)
                token = result["access_token"]
                expiry = time.time() + float(result.get("expires_in", 300))
            except urllib.error.HTTPError as exc:
                try:
                    try:
                        response_body = exc.read(16384)
                    except OSError:
                        response_body = b""
                    message = token_http_error_message(exc.code, response_body)
                finally:
                    exc.close()
                raise DeployError(message) from None
            except urllib.error.URLError as exc:
                raise DeployError(token_network_error_message(exc.reason)) from None
            except OSError as exc:
                raise DeployError(token_network_error_message(exc)) from None
            except (ValueError, KeyError, TypeError):
                raise DeployError("Microsoft Entra 令牌响应无效；请稍后重试或检查网络代理（凭据未输出）") from None
        if not isinstance(token, str) or not token or not math.isfinite(expiry):
            raise DeployError("认证返回无效令牌")
        remember_secret(token)
        self.cache[key] = token, expiry
        return token


class ArmClient:
    """只负责 REST、重试和长任务等待，不掺杂业务计划。"""
    def __init__(self, tokens: TokenProvider, subscription: str, timeout: float = 900,
                 poll_interval: float = 5, retries: int = 4):
        self.tokens = tokens
        self.subscription = subscription
        self.timeout = timeout
        self.poll_interval = poll_interval
        self.retries = retries
        self.opener = urllib.request.build_opener(NoRedirect(), urllib.request.HTTPSHandler(context=ssl.create_default_context()))

    def url(self, path: str, api: str = COGNITIVE_API) -> str:
        """分页/异步地址也必须属于本订阅的 ARM 域，避免令牌泄漏。"""
        url = path if path.startswith("https://") else ARM + urllib.parse.quote(path, safe="/%?=&")
        parsed = urllib.parse.urlsplit(url)
        if parsed.scheme != "https" or parsed.netloc.lower() != "management.azure.com" or parsed.fragment:
            raise DeployError("拒绝访问不可信的 ARM 响应地址")
        decoded = urllib.parse.unquote(parsed.path)
        subscription_root = f"/subscriptions/{self.subscription}".casefold()
        if (decoded.casefold() != subscription_root and not decoded.casefold().startswith(subscription_root + "/")) or ".." in decoded.split("/"):
            raise DeployError("拒绝访问跨订阅的 ARM 响应地址")
        if "api-version" not in urllib.parse.parse_qs(parsed.query):
            url += ("&" if parsed.query else "?") + "api-version=" + api
        return url

    def request(self, method: str, path: str, body: dict | None = None,
                api: str = COGNITIVE_API, deadline: float | None = None) -> tuple[dict, dict, int]:
        url = self.url(path, api)
        payload = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
        refreshed = False
        writing = method.upper() in {"PUT", "POST", "PATCH", "DELETE"}
        for attempt in range(self.retries + 2):
            if deadline is not None and time.monotonic() >= deadline:
                raise DeployError("ARM 操作等待超时；资源可能仍在部署，请在门户确认后重试")
            token = self.tokens.get(self.subscription)
            request = urllib.request.Request(url, data=payload, method=method,
                                             headers={"Authorization": "Bearer " + token,
                                                      "Content-Type": "application/json", "Accept": "application/json"})
            retry_delay = min(2 ** attempt + random.random(), 30)
            request_timeout = 60 if deadline is None else max(0.1, min(60, deadline - time.monotonic()))
            try:
                with self.opener.open(request, timeout=request_timeout) as response:
                    raw = response.read()
                    data = json.loads(raw) if raw else {}
                    if not isinstance(data, dict):
                        error_type = OperationUnknown if writing else DeployError
                        raise error_type("ARM 返回非对象 JSON；写请求的结果需人工核实" if writing else "ARM 返回非对象 JSON")
                    return data, {k.lower(): v for k, v in response.headers.items()}, response.status
            except urllib.error.HTTPError as exc:
                try:
                    error_body = json.loads(exc.read())
                except (ValueError, OSError):
                    error_body = {}
                finally:
                    exc.close()
                error = error_body.get("error", {}) if isinstance(error_body, dict) else {}
                if not isinstance(error, dict):
                    error = {}
                if exc.code == 401 and not refreshed:
                    self.tokens.get(self.subscription, refresh=True)
                    refreshed = True
                    continue
                # 写请求的 5xx 不能证明服务端未执行，不自动再次提交。
                # 明确拒绝的 429 仍按 Retry-After 退避；只读请求保留原重试策略。
                if writing and exc.code >= 500:
                    raise OperationUnknown(f"ARM 写请求返回 HTTP {exc.code}；结果待核实，不会自动重试: {path}") from None
                if exc.code not in RETRY_STATUSES or attempt >= self.retries:
                    raise ArmError(exc.code, text(error.get("code")), redact(text(error.get("message")) or "ARM 请求失败")) from None
                try:
                    retry_delay = max(retry_delay, min(float(exc.headers.get("Retry-After", 0)), 120))
                except (ValueError, TypeError):
                    pass
            except (urllib.error.URLError, TimeoutError, OSError) as exc:
                if writing:
                    raise OperationUnknown(f"ARM 写请求通信中断；结果待核实，不会自动重试: {path}；{redact(str(exc))}") from None
                if attempt >= self.retries:
                    raise DeployError(f"ARM 网络请求失败: {redact(str(exc))}") from None
            except ValueError:
                error_type = OperationUnknown if writing else DeployError
                raise error_type("ARM 写请求返回无效 JSON；结果待核实，不会自动重试" if writing else "ARM 返回了无效 JSON") from None
            except DeployError:
                raise
            except Exception as exc:
                error_type = OperationUnknown if writing else DeployError
                raise error_type(f"ARM {'写请求结果待核实，不会自动重试' if writing else '响应读取失败'}: {path}；{redact(str(exc))}") from None
            if deadline is not None:
                retry_delay = min(retry_delay, max(0, deadline - time.monotonic()))
            time.sleep(retry_delay)
        raise DeployError("ARM 请求重试次数耗尽")

    def get(self, path: str, api: str = COGNITIVE_API) -> dict | None:
        try:
            return self.request("GET", path, api=api)[0]
        except ArmError as exc:
            # 仅 404 代表不存在，不能把权限/网络错误伪装成“需要新建”。
            if exc.status == 404:
                return None
            raise

    def list(self, path: str, api: str = COGNITIVE_API) -> list[dict]:
        """读取 ARM 分页列表；调用方可为资源组等资源管理接口指定 API 版本。"""
        values: list[dict] = []
        seen: set[str] = set()
        while path:
            url = self.url(path, api)
            if url in seen:
                raise DeployError("ARM 分页 nextLink 出现循环")
            seen.add(url)
            result = self.request("GET", url, api=api)[0]
            page = result.get("value", [])
            if not isinstance(page, list):
                raise DeployError("ARM 列表响应格式错误")
            values.extend(page)
            path = result.get("nextLink") or ""
        return values

    def wait_resource(self, path: str, api: str = COGNITIVE_API,
                      deadline: float | None = None, allow_stateless: bool = False) -> dict:
        deadline = deadline if deadline is not None else time.monotonic() + self.timeout
        while time.monotonic() < deadline:
            try:
                resource = self.request("GET", path, api=api, deadline=deadline)[0]
            except ArmError as exc:
                if exc.status != 404:
                    raise OperationUnknown(f"无法确认资源最终状态: {path}；{redact(str(exc))}") from None
                resource = {}
            except DeployError as exc:
                raise OperationUnknown(f"无法确认资源最终状态: {path}；{redact(str(exc))}") from None
            properties = resource.get("properties") or {}
            if not isinstance(properties, dict):
                raise OperationUnknown(f"资源状态响应无效，请核实: {path}")
            state = text(properties.get("provisioningState")).casefold()
            if state == "succeeded" or (resource and not state and allow_stateless):
                return resource
            if state in FAILED_STATES:
                raise OperationFailed(f"资源进入终止状态 {state}: {path}")
            time.sleep(min(self.poll_interval, max(0, deadline - time.monotonic())))
        raise OperationUnknown(f"等待资源超时: {path}；Azure 可能仍在处理，请先核实，不会自动重试或删除")

    def put(self, path: str, body: dict, api: str = COGNITIVE_API, *, allow_stateless: bool = False) -> dict:
        deadline = time.monotonic() + self.timeout
        _, headers, _ = self.request("PUT", path, body, api, deadline)
        # HTTP 202 不是成功；优先等待 Azure 的异步操作状态，再读资源本身。
        try:
            operation = headers.get("azure-asyncoperation") or headers.get("operation-location") or headers.get("location")
            if operation:
                while time.monotonic() < deadline:
                    result, _, status = self.request("GET", operation, api=api, deadline=deadline)
                    state = text(result.get("status") or result.get("properties", {}).get("provisioningState")).casefold()
                    if state in FAILED_STATES:
                        raise OperationFailed("Azure 异步操作失败: " + redact(json.dumps(result.get("error", {"status": state}), ensure_ascii=False)))
                    if state == "succeeded" or (status in (200, 204) and not state):
                        break
                    time.sleep(min(self.poll_interval, max(0, deadline - time.monotonic())))
            return self.wait_resource(path, api, deadline, allow_stateless=allow_stateless or api == RESOURCE_API)
        except (OperationFailed, OperationUnknown):
            raise
        except Exception as exc:
            # 包括轮询权限不足、断网及无效响应。它们不等于部署失败。
            raise OperationUnknown(f"Azure 已接收写请求，但最终结果未确认: {path}；{redact(str(exc))}") from None

# -------------------- 业务层：创建资源、模型选择、配额检查 --------------------
def model_deployment_config(model: dict, sku: dict) -> dict:
    """只比较部署身份、生命周期及所选 SKU，不比较模型展示元数据。"""
    # 模型 PUT 使用 format/name/version；生命周期用于排除 deprecated。
    # SKU 保留容量、配额池、速率及未知配置，避免把真正的 SKU 冲突误合并。
    display_fields = {"description", "displayName", "localizedName", "localizedDescription"}
    config = {key: value for key, value in sku.items() if key not in display_fields}
    config["name"] = text(sku.get("name"))
    config["usageName"] = text(sku.get("usageName")).casefold()
    capacity = sku.get("capacity")
    if capacity is None or isinstance(capacity, dict):
        limits = dict(capacity or {})
        limits.setdefault("minimum", 1)
        if isinstance(limits.get("allowedValues"), list):
            limits["allowedValues"] = sorted(limits["allowedValues"], key=lambda v: json.dumps(v, sort_keys=True))
        config["capacity"] = limits
    return {
        "model": {"name": text(model.get("name")), "format": text(model.get("format")).casefold(),
                  "version": text(model.get("version")),
                  "lifecycleStatus": text(model.get("lifecycleStatus")).casefold()},
        "sku": config,
    }


def model_config_conflicts(configs: list[dict]) -> str:
    """只输出参与选择的配置差异；限制长度，不输出完整目录或凭据。"""
    def flatten(value: dict, prefix: str = "") -> dict:
        result = {}
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else key
            if isinstance(child, dict) and child:
                result.update(flatten(child, path))
            else:
                result[path] = child
        return result
    records = [flatten(config) for config in configs]
    differences = []
    for field in sorted({key for record in records for key in record}):
        values = sorted({json.dumps(record[field], sort_keys=True, ensure_ascii=False)
                         if field in record else "(缺失)" for record in records})
        if len(values) > 1:
            if any(word in field.casefold() for word in ("secret", "password", "token", "key")):
                detail = "(值已隐藏)"
            else:
                detail = " / ".join(value[:160] for value in values[:4])
            differences.append(f"{field}: {detail}")
    return "；".join(differences[:8]) or "目录配置不同"


def ordinary_model_sku_purpose(model: dict, sku: dict) -> tuple[str, str]:
    """为普通 OpenAI 模型区分推理 SKU 与微调用途；不猜测配额池。"""
    if text(model.get("format")).casefold() != "openai":
        return "inference", ""
    usage = text(sku.get("usageName"))
    # 对截图中明确的 -finetune 配额池做边界匹配，不能按模型名随意猜测。
    # 模型 capabilities.fineTune 仅表示支持微调，绝不能据此排除普通推理。
    if re.search(r"(?:^|[.\-_])(?:finetune|fine[-_]?tuning|fine[-_]tune)(?:$|[.\-_])", usage, re.I):
        return "finetune", f"usageName={usage} 明确属于微调用途"
    costs = sku.get("costs") or []
    if isinstance(costs, list):
        hourly_hosting = any(
            isinstance(cost, dict)
            and text(cost.get("name")).casefold() == "hosting"
            and re.search(r"\bhours?\b", text(cost.get("unit")), re.I)
            for cost in costs
        )
        # Hosting 按小时计费是用途线索，但单独出现不足以安全判为微调。
        # Provisioned SKU 本身也可能按小时计费；仅按量普通模型做歧义检查。
        if hourly_hosting and text(sku.get("name")) in {"Standard", "GlobalStandard", "DataZoneStandard"}:
            return "ambiguous", "存在 Hosting 按小时计费，但缺少明确微调配额标识"
    return "inference", ""


def choose_model(catalog: list[dict], item: DeployItem) -> tuple[dict, dict]:
    """空版本使用 Azure 标记的默认版本，不用字符串排序猜“最新版”。"""
    candidates: list[tuple[dict, dict]] = []
    excluded_finetune = 0
    ambiguous_purposes: list[str] = []
    for model in catalog:
        if model.get("name") != item.model_name or text(model.get("format")).casefold() != item.model_format.casefold():
            continue
        if text(model.get("lifecycleStatus")).casefold() == "deprecated":
            continue
        if item.model_version and text(model.get("version")) != item.model_version:
            continue
        for sku in model.get("skus") or []:
            if sku.get("name") == item.sku_name:
                purpose, reason = ordinary_model_sku_purpose(model, sku)
                if purpose == "finetune":
                    excluded_finetune += 1
                    continue
                if purpose == "ambiguous":
                    ambiguous_purposes.append(f"版本 {text(model.get('version'))}: {reason}")
                    continue
                candidates.append((model, sku))
    if ambiguous_purposes:
        raise DeployError(
            f"普通模型 {item.model_name} / SKU {item.sku_name} 存在用途不明确的目录记录："
            + "；".join(sorted(set(ambiguous_purposes))[:4])
            + "；已停止部署，不会猜测或混用微调托管配置"
        )
    if not candidates:
        if excluded_finetune:
            raise DeployError(
                f"普通模型 {item.model_name} {item.model_version or '(默认版本)'} / {item.sku_name} "
                f"仅匹配到 {excluded_finetune} 条微调用途记录，已排除；当前目录没有匹配的普通推理 SKU"
            )
        raise DeployError(f"区域/账号未提供可用组合：{item.model_name} {item.model_version or '(默认版本)'} / {item.sku_name}")
    # 先按 Azure 的默认标记选候选，再去重；显式版本不受默认标记差异影响。
    defaults = [pair for pair in candidates if pair[0].get("isDefaultVersion") is True]
    selected_candidates = defaults if not item.model_version and defaults else candidates
    unique: dict[str, tuple[dict, dict]] = {}
    for model, sku in selected_candidates:
        signature = json.dumps(model_deployment_config(model, sku), sort_keys=True, ensure_ascii=False)
        # 相同签名的部署相关配置一致，因此可使用任一等价记录，而非任意取冲突记录。
        unique.setdefault(signature, (model, sku))
    selected = list(unique.values())
    if len(selected) != 1:
        versions = sorted({text(model.get("version")) or "(缺失版本)" for model, _ in selected})
        if item.model_version or len(versions) == 1:
            # 用户已经指定版本时，不再错误要求重复填写 model_version。
            raise DeployError(
                f"模型 {item.model_name} / 版本 {', '.join(versions)} / SKU {item.sku_name} "
                f"存在 {len(selected)} 条部署配置不同的目录记录；冲突字段："
                + model_config_conflicts([model_deployment_config(model, sku) for model, sku in selected])
                + "；已停止部署，不会随意选择记录。请刷新模型目录或核查 Azure 返回配置"
            )
        raise DeployError(
            f"模型 {item.model_name} 没有唯一默认版本；可选版本：{', '.join(versions)}；"
            "请在模型版本选项中选择明确版本"
        )
    model, sku = selected[0]
    if not text(model.get("version")):
        raise DeployError("Azure 模型目录未返回版本，拒绝猜测")
    limits = sku.get("capacity") or {}
    minimum = limits.get("minimum", 1)
    maximum = limits.get("maximum")
    step = limits.get("step")
    allowed = limits.get("allowedValues")
    if item.capacity < minimum or (maximum is not None and item.capacity > maximum):
        raise DeployError(f"capacity={item.capacity} 超出 SKU 范围 [{minimum}, {maximum}]")
    if step and (item.capacity - minimum) % step:
        raise DeployError(f"capacity 必须从 {minimum} 按步长 {step} 增加")
    if allowed and item.capacity not in allowed:
        raise DeployError(f"capacity 必须属于 {allowed}")
    return model, sku


def same_deployment(existing: dict, item: DeployItem, version: str | None = None) -> bool:
    model = existing.get("properties", {}).get("model", {})
    sku = existing.get("sku", {})
    return (model.get("name") == item.model_name
            and text(model.get("format")).casefold() == item.model_format.casefold()
            and (not (version or item.model_version) or text(model.get("version")) == (version or item.model_version))
            and sku.get("name") == item.sku_name and sku.get("capacity") == item.capacity
            and (not item.rai_policy_name or text(existing.get("properties", {}).get("raiPolicyName")) == item.rai_policy_name))


def quota_details(usages: list[dict], usage_name: str) -> tuple[float | None, str]:
    """精确匹配 usageName；等价重复记录可合并，冲突或异常不得当作零配额。"""
    if not usage_name:
        return None, "模型目录未返回配额标识 usageName"
    matches = []
    for usage in usages:
        if not isinstance(usage, dict):
            continue
        name = usage.get("name")
        value = name.get("value") if isinstance(name, dict) else name if isinstance(name, str) else ""
        # 不使用 localizedValue 或模型名猜测配额池。
        if text(value).casefold() == usage_name.casefold():
            matches.append(usage)
    if not matches:
        return None, f"配额接口返回 {len(usages)} 条记录，但没有匹配 usageName={usage_name} 的记录"
    balances: set[tuple[float, float, str]] = set()
    for usage in matches:
        try:
            if isinstance(usage.get("limit"), bool) or isinstance(usage.get("currentValue"), bool):
                raise ValueError("boolean quota")
            limit, current = float(usage["limit"]), float(usage["currentValue"])
            if not math.isfinite(limit) or not math.isfinite(current) or limit < 0 or current < 0:
                raise ValueError("invalid quota")
            balances.add((limit, current, text(usage.get("unit")).casefold()))
        except (KeyError, TypeError, ValueError, OverflowError):
            return None, f"配额记录 {usage_name} 的 limit/currentValue 缺失或异常"
    if len(balances) != 1:
        return None, f"配额标识 {usage_name} 有 {len(matches)} 条记录，其上限、已用量或单位不一致"
    limit, current, _ = next(iter(balances))
    return max(0, limit - current), ""


def quota_remaining(usages: list[dict], usage_name: str) -> float | None:
    # 保留原函数签名供其它调用方使用。
    return quota_details(usages, usage_name)[0]


class Deployer:
    def __init__(self, client: ArmClient, update_existing: bool = False,
                 skip_quota_check: bool = False, register_provider: bool = False):
        self.client = client
        self.update_existing = update_existing
        self.skip_quota_check = skip_quota_check
        self.register_provider = register_provider
        self.ready: dict[str, dict] = {}
        self.unresolved: dict[str, str] = {}  # 共用资源状态未知时，后续模型不重复创建。
        self.catalogs: dict[str, list[dict]] = {}
        # 共享资源和模型目录串行准备；模型 PUT / 轮询不持有此锁。
        self.state_lock = RLock()
        # 同订阅不同资源组也共享配额余额，预留必须原子执行。
        self.quota_lock = RLock()
        # 失败也不归还余额，避免 Azure 已提交但本地超时导致重复占用。
        self.quota_budget: dict[tuple[str, str], float] = {}
        self.provider_ready = False

    def ensure_provider(self) -> None:
        if self.provider_ready:
            return
        path = f"/subscriptions/{self.client.subscription}/providers/Microsoft.CognitiveServices"
        state = self.client.request("GET", path, api=PROVIDER_API)[0].get("registrationState")
        if state != "Registered":
            if not self.register_provider:
                raise DeployError("Microsoft.CognitiveServices 尚未注册；先注册或显式添加 --register-provider")
            self.client.request("POST", path + "/register", {}, PROVIDER_API)
            deadline = time.monotonic() + self.client.timeout
            while state != "Registered":
                if time.monotonic() >= deadline:
                    raise DeployError("资源提供程序注册超时")
                time.sleep(min(self.client.poll_interval, max(0, deadline - time.monotonic())))
                state = self.client.request("GET", path, api=PROVIDER_API, deadline=deadline)[0].get("registrationState")
        self.provider_ready = True

    def ensure(self, path: str, body: dict, api: str = COGNITIVE_API) -> dict:
        key = path.casefold()
        if key in self.unresolved:
            raise OperationUnknown(self.unresolved[key])
        try:
            if key not in self.ready:
                existing = self.client.get(path, api)
                if existing is None:
                    self.ready[key] = self.client.put(path, body, api)
                else:
                    state = text(existing.get("properties", {}).get("provisioningState")).casefold()
                    if state != "succeeded" and not (api == RESOURCE_API and not state):
                        existing = self.client.wait_resource(path, api)
                    self.ready[key] = existing
            return self.ready[key]
        except OperationUnknown as exc:
            self.unresolved[key] = str(exc)
            raise

    def require_existing(self, path: str, label: str, api: str = COGNITIVE_API) -> None:
        """已选已有资源在执行时再次核对；404 绝不降级成创建。"""
        key = path.casefold()
        if key in self.unresolved:
            raise OperationUnknown(self.unresolved[key])
        try:
            value = self.client.get(path, api=api)
            if value is None:
                raise DeployError(f"所选{label}已不存在，请返回部署目标重新选择；不会自动重建")
            state = text(value.get("properties", {}).get("provisioningState")).casefold()
            if state != "succeeded" and not (api == RESOURCE_API and not state):
                value = self.client.wait_resource(path, api)
            self.ready[key] = value
        except OperationUnknown as exc:
            self.unresolved[key] = str(exc)
            raise

    def progress(self, stage: str) -> None:
        callback = getattr(self, "progress_callback", None)
        if callback is not None:
            callback(stage)

    def prepare(self, item: DeployItem) -> tuple[dict, dict]:
        self.ensure_provider()
        self.progress("资源组")
        group_id = f"/subscriptions/{item.subscription_id}/resourceGroups/{item.resource_group}"
        try:
            if item.resource_group_mode == "existing":
                self.require_existing(group_id, "资源组", RESOURCE_API)
            self.ensure(group_id, {"location": item.location}, RESOURCE_API)
        except ArmError as exc:
            # 列表可见不代表具有单个资源组的 read 权限；403 不可当作 404 自动新建。
            if exc.status == 403 and "Microsoft.Resources/subscriptions/resourceGroups/read" in str(exc):
                raise DeployError(
                    f"SP 无权读取订阅 {item.subscription_id} 下的资源组 {item.resource_group} "
                    "(Microsoft.Resources/subscriptions/resourceGroups/read)。"
                    "请确认目标订阅和资源组名称，并让管理员在该资源组或订阅范围授予此 SP "
                    "所需的读取及创建资源权限；权限刚授予时请重新认证后重试。"
                ) from exc
            raise
        self.progress("Foundry")
        if item.account_mode == "existing":
            self.require_existing(item.account_id, "Foundry")
        account = self.ensure(item.account_id, {
            "kind": "AIServices", "location": item.location, "sku": {"name": "S0"},
            "identity": {"type": "SystemAssigned"},
            "properties": {"customSubDomainName": item.account_name, "allowProjectManagement": True},
        })
        props = account.get("properties", {})
        if account.get("kind") != "AIServices" or normalize_location(account.get("location", "")) != item.location:
            raise DeployError("已有账号类型或区域不匹配；不会自动改造/迁移旧资源")
        if not props.get("allowProjectManagement"):
            raise DeployError("已有账号未启用 allowProjectManagement；请人工确认后启用")
        self.progress("Project")
        if item.project_mode == "existing":
            self.require_existing(item.project_id, "Project")
        project = self.ensure(item.project_id, {
            "location": item.location, "identity": {"type": "SystemAssigned"},
            "properties": {"displayName": item.project_display_name or item.project_name,
                           "description": item.project_description},
        })
        if normalize_location(project.get("location", "")) != item.location:
            raise DeployError("已有项目区域不匹配")
        return account, project

    def reserve_quota(self, item: DeployItem, sku: dict, existing: dict | None) -> str:
        with self.quota_lock:
            return self._reserve_quota(item, sku, existing)

    def _reserve_quota(self, item: DeployItem, sku: dict, existing: dict | None) -> str:
        usage = text(sku.get("usageName"))
        # 同型号+SKU 只计算容量差额；换型号/SKU 不假设旧配额可抵扣。
        old = existing or {}
        old_model = old.get("properties", {}).get("model", {})
        old_sku = old.get("sku", {})
        # 版本变更也可能切换配额池；必须按旧版本目录的 usageName 核实，不能只看模型名。
        old_usage_names = {
            text(old_catalog_sku.get("usageName")).casefold()
            for old_catalog in self.catalogs.get(item.account_id.casefold(), [])
            if old_catalog.get("name") == old_model.get("name")
            and text(old_catalog.get("format")).casefold() == text(old_model.get("format")).casefold()
            and text(old_catalog.get("version")) == text(old_model.get("version"))
            for old_catalog_sku in old_catalog.get("skus") or []
            if old_catalog_sku.get("name") == old_sku.get("name")
        }
        same_pool = bool(usage) and old_usage_names == {usage.casefold()}
        increase = max(0, item.capacity - int(old_sku.get("capacity") or 0)) if same_pool else item.capacity
        if not increase:
            return "不增加容量；未释放本次运行的保守余额"
        # 部分当前 Foundry 模型目录不会返回 usageName。没有稳定映射时不能用模型名猜配额池，
        # 因此不阻断真实部署，交由 Azure 创建接口按当前订阅配额做最终校验。
        if not usage:
            return "模型目录未返回配额标识，跳过本地配额预检，交由 Azure 最终校验"
        try:
            usages = self.client.list(
                f"/subscriptions/{item.subscription_id}/providers/Microsoft.CognitiveServices/locations/{item.location}/usages",
                api=USAGE_API,
            )
        except ArmError as exc:
            if self.skip_quota_check and exc.status in {401, 403, 404}:
                return f"配额接口不可读取（HTTP {exc.status}），按选项交由 Azure 最终校验"
            if exc.status in {401, 403}:
                raise DeployError(
                    f"配额接口访问被拒绝（HTTP {exc.status}）；请核查认证及订阅级配额读取权限 "
                    "(Cognitive Services Usages Reader)。这不代表配额为零。"
                    "可显式勾选“配额读取失败时仍尝试部署”交由 Azure 最终校验"
                ) from exc
            raise
        remaining, reason = quota_details(usages, usage)
        if remaining is None:
            if self.skip_quota_check:
                return reason + "；按选项交由 Azure 最终校验"
            raise DeployError(
                reason + "；无法核实配额，不等同于 SP 无效或配额为零。"
                "可显式勾选“配额读取失败时仍尝试部署”交由 Azure 创建接口校验"
            )
        key = item.location, usage.casefold()
        available = min(remaining, self.quota_budget.get(key, remaining))
        if increase > available:
            raise DeployError(f"配额不足：{usage} 还需 {increase} capacity，可用 {available:g}；不会回收其他部署")
        # 失败也不乐观归还：超时请求可能已经在 Azure 成功。
        self.quota_budget[key] = available - increase
        return f"{usage}: 本次新增 {increase}，保守剩余 {available - increase:g}"

    def ensure_guard_policy(self, item: DeployItem) -> None:
        """只创建本控制台的固定、阻止型策略；已有同名不同配置一律拒绝覆盖。"""
        threshold = GUARD_POLICIES[item.rai_policy_name]
        path = item.account_id + "/raiPolicies/" + item.rai_policy_name
        filters = [{"name": category, "source": source, "severityThreshold": threshold,
                    "enabled": True, "blocking": True}
                   for source in ("Prompt", "Completion")
                   for category in ("Hate", "Sexual", "Violence", "SelfHarm")]
        if path.casefold() in self.unresolved:
            raise OperationUnknown(self.unresolved[path.casefold()])
        policy = self.client.get(path)
        if policy is None:
            body = {"properties": {"basePolicyName": "Microsoft.Default",
                                   "mode": "Blocking", "contentFilters": filters}}
            # raiPolicies 的成功响应不一定含 provisioningState；允许轮询无状态的资源。
            try:
                policy = self.client.put(path, body, allow_stateless=True)
            except OperationUnknown as exc:
                self.unresolved[path.casefold()] = str(exc)
                raise
        properties = policy.get("properties") or {}
        # 不复用显式异步/批注策略，否则界面显示“拦截”而 Azure 实际没有阻断。
        if properties.get("mode", "Blocking") != "Blocking":
            raise DeployError("同名内容防护策略不是阻止模式；不会覆盖，请在 Azure 检查策略")
        actual = properties.get("contentFilters") or []
        expected = {(f["source"].casefold(), f["name"].casefold()): f for f in filters}
        found = {(str(f.get("source", "")).casefold(), str(f.get("name", "")).casefold()): f
                 for f in actual if isinstance(f, dict)}
        if any(key not in found or found[key].get("severityThreshold") != value["severityThreshold"]
               or found[key].get("blocking") is not True or found[key].get("enabled") is not True
               for key, value in expected.items()):
            raise DeployError("同名内容防护策略与本控制台预设不符；不会覆盖，请在 Azure 检查策略")

    def deploy(self, item: DeployItem, result: dict) -> None:
        if item.rai_policy_name and item.rai_policy_name not in GUARD_POLICIES:
            # 用户手动指定的旧策略仍需预先存在，且不会被本工具更改。
            policy_path = item.account_id + "/raiPolicies/" + urllib.parse.quote(item.rai_policy_name, safe="")
            if self.client.get(policy_path) is None:
                raise DeployError(f"目标 Foundry 资源中不存在内容筛选器 {item.rai_policy_name}；请重新读取并选择")
        # 共用资源准备完毕后释放锁，让不同模型同时提交、等待。
        with self.state_lock:
            account, project = self.prepare(item)
            if item.rai_policy_name in GUARD_POLICIES:
                self.ensure_guard_policy(item)
        result["account_endpoint"] = account.get("properties", {}).get("endpoint", "")
        result["project_endpoints"] = project.get("properties", {}).get("endpoints") or {}
        existing = self.client.get(item.deployment_id)
        if existing:
            state = text(existing.get("properties", {}).get("provisioningState")).casefold()
            if state != "succeeded":
                existing = self.client.wait_resource(item.deployment_id)
            # 版本留空意味着保留已有版本，不会每次重跑都升级到新默认版本。
            if same_deployment(existing, item):
                result.update(status="unchanged", resolved_model_version=existing.get("properties", {}).get("model", {}).get("version", ""))
                return
            if not self.update_existing:
                raise DeployError("同名部署配置不同；默认不覆盖。确认后添加 --update-existing")
        self.progress("核对模型")
        catalog_key = item.account_id.casefold()
        with self.state_lock:
            if catalog_key not in self.catalogs:
                self.catalogs[catalog_key] = self.client.list(item.account_id + "/models")
        selection = item
        if existing and not item.model_version:
            old_model = existing.get("properties", {}).get("model", {})
            if old_model.get("name") == item.model_name and text(old_model.get("format")).casefold() == item.model_format.casefold():
                selection = replace(item, model_version=text(old_model.get("version")))
        model, sku = choose_model(self.catalogs[catalog_key], selection)
        version = text(model["version"])
        result["resolved_model_version"] = version
        props: dict[str, Any] = {"model": {"format": item.model_format, "name": item.model_name, "version": version}}
        if existing:
            previous = existing.get("properties", {})
            # PUT 保留内容安全和版本升级等可写配置，不静默重置策略。
            for key in ("raiPolicyName", "versionUpgradeOption", "dynamicThrottlingEnabled",
                        "parentDeploymentName", "spilloverDeploymentName"):
                if key in previous:
                    props[key] = previous[key]
            if previous.get("scaleSettings") or previous.get("capacitySettings"):
                raise DeployError("已有部署含 scaleSettings/内部 capacitySettings；请人工处理，脚本不做隐式转换")
        if item.rai_policy_name:
            props["raiPolicyName"] = item.rai_policy_name
        body: dict[str, Any] = {"sku": {"name": item.sku_name, "capacity": item.capacity}, "properties": props}

        if existing:
            if "tags" in existing:
                body["tags"] = existing["tags"]
            if existing.get("sku", {}).get("name") == item.sku_name:
                for key in ("tier", "size", "family"):
                    if key in existing["sku"]:
                        body["sku"][key] = existing["sku"][key]
        self.progress("核对配额")
        result["quota_check"] = self.reserve_quota(item, sku, existing)
        self.progress("提交模型并等待 Azure")
        deployed = self.client.put(item.deployment_id, body)
        if not same_deployment(deployed, item, version):
            raise OperationUnknown("Azure 返回成功状态，但最终部署配置与计划不一致，请人工核实；不会自动重试")
        result["status"] = "updated" if existing else "created"
        result["rai_policy"] = text(deployed.get("properties", {}).get("raiPolicyName"))

# -------------------- 输出与命令行 --------------------
def result_row(item: DeployItem, status: str) -> dict:
    return {**asdict(item), "status": status, "resolved_model_version": "",
            "account_resource_id": item.account_id, "project_resource_id": item.project_id,
            "deployment_resource_id": item.deployment_id, "account_endpoint": "", "project_endpoints": {},
            "quota_check": "", "error": "", "subscription_name": "", "tpm": "", "endpoint": "",
            "azure_openai_endpoint": "", "rai_policy": "", "streaming_filter_mode": "", "report_warnings": []}


def csv_safe(value: Any) -> Any:
    """CSV 用 Excel 打开时，不把描述或错误文本当成公式执行。"""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def write_report(directory: Path, rows: list[dict], mode: str, duplicates: int, *, exporter=None) -> None:
    # 白名单结果字段中也可能含服务端回显；递归脱敏，再写 JSON 与 CSV。
    def scrub(value):
        if isinstance(value, str):
            return redact(value, limit=None)
        if isinstance(value, dict):
            return {redact(str(key), limit=None): scrub(child) for key, child in value.items()}
        if isinstance(value, list):
            return [scrub(child) for child in value]
        return value
    rows = scrub(rows)
    summary: dict[str, int] = {}
    for row in rows:
        summary[row["status"]] = summary.get(row["status"], 0) + 1
    report = {"mode": mode, "updated_at_utc": datetime.now(timezone.utc).isoformat(),
              "duplicates_ignored": duplicates, "summary": summary, "deployments": rows}
    # 每条部署后落盘，单文件原子替换，避免进程中断留下半截 JSON。
    temporary = directory / "results.json.tmp"
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(directory / "results.json")
    temporary = directory / "results.csv.tmp"
    with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(dict.fromkeys(key for row in rows for key in row)))
        writer.writeheader()
        writer.writerows({k: csv_safe(v) for k, v in row.items()} for row in rows)
    temporary.replace(directory / "results.csv")

    # 将同一份已脱敏数据输出为截图风格的 Excel；打开中的 Excel 可能被 Windows 锁定。
    # 仅 Excel 写入失败时，保留部署流程和 CSV/JSON，并明确标记警告，不误报云端失败。
    try:
        write_excel_report(directory / "results.xlsx", report)
    except OSError:
        warning = "Excel 报告无法更新（文件可能正被打开或目录不可写）；请以 results.json / results.csv 为准，关闭 Excel 后再试。"
        report["report_warnings"] = [warning]
        print("警告：" + warning, file=sys.stderr, flush=True)
        temporary = directory / "results.json.tmp"
        temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(directory / "results.json")

    if exporter is not None:
        exporter.write(report)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Azure AI Foundry 部署工具：默认离线预览，--apply 才部署。")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--config", type=Path, help="JSON 或 CSV 部署计划")
    source.add_argument("--excel", type=Path, help="兼容原 Excel 工作簿（不改写原文件）")
    parser.add_argument("--sheet", help="Excel 工作表名；默认活动工作表")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--apply", action="store_true", help="实际创建云资源，可能产生费用")
    mode.add_argument("--dry-run", action="store_true", help="显式预览，与默认行为相同；不认证、不请求 Azure")
    parser.add_argument("--auth", choices=["cli", "service-principal"], default=None, help="认证方式：JSON 有 credentials 时默认服务主体，否则复用 az login")
    parser.add_argument("--update-existing", action="store_true", help="允许更新配置不同的同名模型部署（不改变账号/项目配置）")
    parser.add_argument("--register-provider", action="store_true", help="允许注册 Microsoft.CognitiveServices 提供程序")
    parser.add_argument("--skip-quota-check", action="store_true", help="配额信息缺失时允许继续；已知不足仍失败")
    parser.add_argument("--detailed-report", action="store_true", help="部署成功后读取订阅、端点、TPM 和策略详情")
    parser.add_argument("--max-concurrency", type=int, default=3, help="模型任务并发数 1–8；共享资源先准备，模型提交及等待并行")
    parser.add_argument("--subscription-id", help="只运行指定订阅的计划，不修改计划中的订阅")
    parser.add_argument("--limit", type=int, help="只运行前 N 条，便于先做小范围试部署")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "output", help="报告父目录；每次新建独立子目录")
    parser.add_argument("--poll-timeout", type=float, default=900, help="单资源最长等待秒数，默认 900")
    parser.add_argument("--poll-interval", type=float, default=5, help="轮询间隔秒数，默认 5")
    parser.add_argument("--retries", type=int, default=4, help="ARM 临时错误最大重试次数，默认 4")
    args = parser.parse_args(argv)
    if not 1 <= args.max_concurrency <= 8:
        parser.error("--max-concurrency 必须为 1–8")
    if args.limit is not None and args.limit < 1:
        parser.error("--limit 必须大于 0")
    if not all(math.isfinite(v) and v > 0 for v in (args.poll_timeout, args.poll_interval)) or args.retries < 0:
        parser.error("超时/轮询间隔必须为有限正数，重试次数不能为负数")
    return args


def main(argv: list[str] | None = None, *, config_document: Any = _MISSING,
         run_directory: Path | None = None, report_exporter=None) -> int:
    args = parse_args(argv)
    directory = run_directory
    def save_report(rows, mode, duplicates):
        write_report(directory, rows, mode, duplicates, exporter=report_exporter)

    rows: list[dict] = []
    mode = "apply" if args.apply else "dry-run"
    try:
        source_path = args.config or args.excel
        document = config_document
        if source_path.suffix.lower() == ".json" and document is _MISSING:
            document = read_json_document(source_path)
        plan, duplicates = load_plan(source_path, args.sheet, document=document)
        credentials = document.get("credentials") if isinstance(document, dict) else None
        if args.subscription_id:
            try:
                subscription = str(uuid.UUID(args.subscription_id))
            except ValueError:
                raise PlanError("--subscription-id 必须是 UUID") from None
            plan = [item for item in plan if item.subscription_id == subscription]
        if args.limit:
            plan = plan[:args.limit]
        if not plan:
            raise PlanError("筛选后计划为空")
        if args.apply and any(uuid.UUID(item.subscription_id).int == 0 for item in plan):
            raise PlanError("请把示例中的全零 subscription_id 替换为真实订阅 ID")
        run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f") + "_" + uuid.uuid4().hex[:6]
        if directory is None:
            directory = args.output_dir.resolve() / run_id
            directory.mkdir(parents=True, exist_ok=False)
        else:
            directory.mkdir(parents=True, exist_ok=True)
        mode = "apply" if args.apply else "dry-run"
        rows = [result_row(item, "pending" if args.apply else "planned") for item in plan]
        save_report(rows, mode, duplicates)
        print(f"模式={mode}，部署数={len(plan)}，已去重={duplicates}")
        print(f"报告目录：{directory}")
        if not args.apply:
            for item in plan[:20]:
                print(f"  {item.subscription_id} / {item.location} / {item.account_name} / {item.deployment_name}: "
                      f"{item.model_name}@{item.model_version or 'Azure默认版本/保留已有版本'} {item.sku_name} capacity={item.capacity}")
            if len(plan) > 20:
                print("  其余条目请查看报告。")
            print("离线预览完成。未认证、未请求 Azure；模型可用性和配额将在 --apply 时核实。")
            return 0
        auth_mode = args.auth or ("service-principal" if credentials is not None else "cli")
        tokens = TokenProvider(auth_mode, credentials=credentials) if credentials is not None else TokenProvider(auth_mode)
        # 在首个写请求前验证各订阅认证；服务主体模式一次运行只使用一个租户/应用。
        # Preflight each subscription independently: one inaccessible target must not
        # prevent authorized subscriptions from proceeding with their own deployments.
        subscriptions = sorted({item.subscription_id for item in plan})
        ready_subscriptions: set[str] = set()
        for subscription in subscriptions:
            try:
                tokens.get(subscription)
            except DeployError as exc:
                for row in rows:
                    if row["subscription_id"] == subscription:
                        row.update(status="auth_failed", error=redact(str(exc)))
                print(f"订阅 {subscription}：SP 认证失败", flush=True)
            else:
                ready_subscriptions.add(subscription)
        # Provider registration is subscription-scoped. Complete this phase before
        # starting workers so concurrent resource groups do not race to register it.
        for subscription in sorted(ready_subscriptions):
            preflight = Deployer(ArmClient(tokens, subscription, args.poll_timeout, args.poll_interval, args.retries),
                                 args.update_existing, args.skip_quota_check, args.register_provider)
            try:
                preflight.ensure_provider()
            except DeployError as exc:
                for row in rows:
                    if row["subscription_id"] == subscription:
                        row.update(status="not_run" if isinstance(exc, OperationUnknown) else "failed",
                                   error=redact(str(exc)) + "；尚未开始模型部署")
                ready_subscriptions.remove(subscription)
                print(f"订阅 {subscription}：Provider 预检失败", flush=True)
        if len(ready_subscriptions) != len(subscriptions):
            save_report(rows, mode, duplicates)
            for index, row in enumerate(rows):
                if row["status"] in {"auth_failed", "failed", "not_run"}:
                    item = plan[index]
                    print(f"[{index + 1}/{len(plan)}] {item.subscription_id}/{item.location}/{item.deployment_name}  {row['status']}: {row['error']}", flush=True)
        if not ready_subscriptions:
            print("没有订阅通过预检，请查看报告中的逐订阅错误", flush=True)
            return 1
        # 调度单位是一条模型任务，而不是资源组；并发不会生成额外计划行。
        tasks = [(index, item) for index, item in enumerate(plan)
                 if item.subscription_id in ready_subscriptions]
        group_states = {}
        quota_states = {}
        for _, item in tasks:
            group_key = (item.subscription_id, item.resource_group.casefold())
            group_states.setdefault(group_key, {"lock": RLock(), "ready": {}, "unresolved": {}, "catalogs": {}})
            quota_states.setdefault(item.subscription_id, {"lock": RLock(), "budget": {}})
        # 详情登记 API Key 与报告导出共享字典，使用同一锁，报告仍仅主线程落盘。
        report_lock = RLock()
        output_lock = RLock()
        queue: Queue[tuple[int, dict]] = Queue()

        def emit_progress(index, stage):
            with output_lock:
                print("__FOUNDRY_PROGRESS__" + json.dumps(
                    {"index": index, "stage": stage}, ensure_ascii=False), flush=True)

        def perform_model(index, item):
            row = result_row(item, "running")
            try:
                # 每个模型独立客户端/令牌缓存，避免并发刷新和回调互相覆盖。
                worker_tokens = (TokenProvider(auth_mode, credentials=credentials)
                                 if credentials is not None else TokenProvider(auth_mode))
                worker_tokens.cache = dict(tokens.cache)
                client = ArmClient(worker_tokens, item.subscription_id,
                                   args.poll_timeout, args.poll_interval, args.retries)
                worker = Deployer(client, args.update_existing,
                                  args.skip_quota_check, args.register_provider)
                worker.provider_ready = True
                shared = group_states[(item.subscription_id, item.resource_group.casefold())]
                worker.state_lock = shared["lock"]
                worker.ready = shared["ready"]
                worker.unresolved = shared["unresolved"]
                worker.catalogs = shared["catalogs"]
                quota = quota_states[item.subscription_id]
                worker.quota_lock = quota["lock"]
                worker.quota_budget = quota["budget"]
                worker.progress_callback = lambda stage: emit_progress(index, stage)
                worker.progress("准备共享资源")
                worker.deploy(item, row)
                worker.progress("读取部署结果")
                if args.detailed_report and report_exporter is not None:
                    try:
                        from foundry_details import DeploymentDetails
                        with report_lock:
                            DeploymentDetails(client, report_exporter).enrich(item, row)
                    except Exception:
                        row.setdefault("report_warnings", []).append(
                            "补充详情读取异常；部署状态保留，请到 Azure 门户核实空白字段。")
            except OperationUnknown as exc:
                row.update(status="interrupted", error=redact(str(exc)))
            except (DeployError, ValueError, TypeError, KeyError) as exc:
                row.update(status="failed", error=redact(str(exc)))
            except Exception:
                row.update(status="interrupted", error="未预期的部署异常，无法确定最终状态；请先核实 Azure 资源，不会自动重试")
            finally:
                # 无论初始化、资源准备或模型部署失败，都必须返回一条结果，避免主线程挂起。
                queue.put((index, row))

        print(f"模型并发上限={args.max_concurrency}；模型任务数={len(tasks)}；资源组数={len(group_states)}", flush=True)
        try:
            with ThreadPoolExecutor(max_workers=args.max_concurrency) as pool:
                futures = [pool.submit(perform_model, index, item) for index, item in tasks]
                for _ in tasks:
                    index, row = queue.get()
                    rows[index] = row
                    with output_lock:
                        print(f"[{index + 1}/{len(plan)}] {plan[index].subscription_id}/{plan[index].location}/{plan[index].deployment_name}  {row['status']}" + (f": {row['error']}" if row.get("error") else ""), flush=True)
                    with report_lock:
                        save_report(rows, mode, duplicates)
                for future in futures:
                    future.result()
        except KeyboardInterrupt:
            for row in rows:
                if row["status"] in ("pending", "running"):
                    row.update(status="interrupted", error="用户中断；Azure 端任务可能仍在运行，请检查资源状态")
            save_report(rows, mode, duplicates)
            return 130
        failures = sum(row["status"] in {"failed", "auth_failed"} for row in rows)
        successes = sum(row["status"] in {"created", "updated", "unchanged"} for row in rows)
        unknown = sum(row["status"] == "interrupted" for row in rows)
        not_run = sum(row["status"] == "not_run" for row in rows)
        print(f"已完成：成功/无需变更 {successes}，失败 {failures}，未执行 {not_run}，待核实 {unknown}；报告目录：{directory}")
        return 1 if failures or unknown or not_run else 0
    except (PlanError, DeployError, OSError, ValueError) as exc:
        if directory is not None and not rows:
            try:
                save_report([{"status": "config_failed", "error": redact(str(exc))}], mode, 0)
            except OSError:
                pass  # 磁盘不可写时无法承诺落盘；仍以非零退出码和控制台报错结束。
        print("错误：" + redact(str(exc)), file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("已中断；若已提交 Azure 请求，请在门户检查状态。", file=sys.stderr)
        return 130


if __name__ == "__main__":
    # 兼容 Windows 中文控制台；日志/报告始终使用 UTF-8。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())

